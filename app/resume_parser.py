"""
Resume parser supporting PDF and DOCX formats.
Enforces size limit, file signature validation, and extracts text locally.
Does not log resume contents.
"""
import io
import fitz  # PyMuPDF
import docx
from typing import Tuple


PDF_MAGIC = b"%PDF-"
DOCX_ZIP_MAGIC = b"PK\x03\x04"


class ResumeExtractionError(Exception):
    """Raised when resume extraction fails due to format, corruption or size."""
    pass


def validate_resume_file(data: bytes, filename: str, max_size_bytes: int) -> str:
    """
    Validate file size and magic signature bytes.
    Returns detected format ('pdf' or 'docx').
    """
    if len(data) == 0:
        raise ResumeExtractionError("File is empty.")
    
    if len(data) > max_size_bytes:
        raise ResumeExtractionError(f"File size exceeds maximum allowed limit of {max_size_bytes // 1024 // 1024}MB.")

    lower_filename = filename.lower()
    
    if data.startswith(PDF_MAGIC) and lower_filename.endswith(".pdf"):
        return "pdf"
    
    if data.startswith(DOCX_ZIP_MAGIC) and lower_filename.endswith(".docx"):
        return "docx"
    
    raise ResumeExtractionError(
        "Invalid file type. Only standard PDF (.pdf) and Microsoft Word (.docx) files are supported."
    )


def extract_text_from_pdf(data: bytes) -> str:
    """Extract text from PDF using PyMuPDF locally."""
    try:
        doc = fitz.open(stream=data, filetype="pdf")
        if doc.page_count == 0:
            raise ResumeExtractionError("PDF file contains no pages.")
        
        extracted_pages = []
        for page_idx in range(doc.page_count):
            page = doc[page_idx]
            text = page.get_text("text")
            if text:
                extracted_pages.append(text.strip())
        
        doc.close()
        full_text = "\n\n".join(extracted_pages).strip()
        if not full_text:
            raise ResumeExtractionError("Could not extract readable text from PDF (it may be scanned or empty).")
        return full_text
    except ResumeExtractionError:
        raise
    except Exception as e:
        raise ResumeExtractionError(f"Failed to process PDF: {e}") from e


def extract_text_from_docx(data: bytes) -> str:
    """Extract text from DOCX using python-docx locally."""
    try:
        stream = io.BytesIO(data)
        document = docx.Document(stream)
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text and cell_text not in paragraphs:
                        paragraphs.append(cell_text)
                        
        full_text = "\n\n".join(paragraphs).strip()
        if not full_text:
            raise ResumeExtractionError("Could not extract readable text from DOCX document.")
        return full_text
    except ResumeExtractionError:
        raise
    except Exception as e:
        raise ResumeExtractionError(f"Failed to process DOCX file: {e}") from e


def extract_resume_text(data: bytes, filename: str, max_size_bytes: int) -> Tuple[str, str]:
    """
    Validate and extract text from uploaded resume bytes.
    Returns (extracted_text, detected_type).
    """
    detected_type = validate_resume_file(data, filename, max_size_bytes)
    if detected_type == "pdf":
        text = extract_text_from_pdf(data)
    elif detected_type == "docx":
        text = extract_text_from_docx(data)
    else:
        raise ResumeExtractionError(f"Unsupported type: {detected_type}")
    
    return text, detected_type
